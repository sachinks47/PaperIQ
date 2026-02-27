import streamlit as st
import pdfplumber
import docx
import re
import numpy as np
import plotly.graph_objects as go
import heapq
import sqlite3
import hashlib
import json
import datetime
from collections import Counter
from textblob import TextBlob
from fpdf import FPDF

# -----------------------------
# Database Setup
# -----------------------------
def init_db():
    conn = sqlite3.connect('paperiq_users.db')
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS users 
                 (username TEXT PRIMARY KEY, gmail TEXT, password TEXT, role TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS history
                 (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT, filename TEXT, date TEXT, data TEXT)''')
    conn.commit()
    conn.close()

def add_user(username, gmail, password, role):
    conn = sqlite3.connect('paperiq_users.db')
    c = conn.cursor()
    hashed_pwd = hashlib.sha256(password.encode()).hexdigest()
    try:
        c.execute("INSERT INTO users VALUES (?,?,?,?)", (username, gmail, hashed_pwd, role))
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    finally:
        conn.close()

def verify_user(username, password):
    conn = sqlite3.connect('paperiq_users.db')
    c = conn.cursor()
    hashed_pwd = hashlib.sha256(password.encode()).hexdigest()
    c.execute("SELECT gmail, role FROM users WHERE username=? AND password=?", (username, hashed_pwd))
    result = c.fetchone()
    conn.close()
    return result

def save_history(username, filename, data):
    conn = sqlite3.connect('paperiq_users.db')
    c = conn.cursor()
    date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    data_json = json.dumps(data)
    c.execute("INSERT INTO history (username, filename, date, data) VALUES (?, ?, ?, ?)", 
              (username, filename, date_str, data_json))
    conn.commit()
    conn.close()

def get_history(username):
    conn = sqlite3.connect('paperiq_users.db')
    c = conn.cursor()
    c.execute("SELECT id, filename, date, data FROM history WHERE username=? ORDER BY id DESC", (username,))
    rows = c.fetchall()
    conn.close()
    return rows

def delete_history(row_id):
    conn = sqlite3.connect('paperiq_users.db')
    c = conn.cursor()
    c.execute("DELETE FROM history WHERE id=?", (row_id,))
    conn.commit()
    conn.close()

# Initialize DB on start
init_db()

# -----------------------------
# Configuration & Styling
# -----------------------------
st.set_page_config(page_title="PaperIQ", page_icon="✨", layout="wide")

# Session State Initialization
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
if 'auth_mode' not in st.session_state:
    st.session_state.auth_mode = "login"
if 'user_data' not in st.session_state:
    st.session_state.user_data = {}
if 'analysis_results' not in st.session_state:
    st.session_state.analysis_results = None
if 'summary_length_slider' not in st.session_state:
    st.session_state.summary_length_slider = "Medium"

# CSS Styling (Teal Green & Animated Dark Mode)
st.markdown("""
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;600;800&display=swap');

        /* Animated Background Effect */
        @keyframes gradientBG {
            0% { background-position: 0% 50%; }
            50% { background-position: 100% 50%; }
            100% { background-position: 0% 50%; }
        }

        [data-testid="stAppViewContainer"] {
            background: linear-gradient(-45deg, #0a0f14, #0d1b22, #081217, #050a10);
            background-size: 400% 400%;
            animation: gradientBG 15s ease infinite;
            color: #ffffff;
            font-family: 'Inter', sans-serif;
        }
        
        [data-testid="stSidebar"] {
            background-color: rgba(10, 15, 20, 0.95) !important;
            border-right: 1px solid rgba(45, 212, 191, 0.1);
        }

        /* Hero Text Styling */
        .hero-title { 
            text-align: center; 
            margin-top: 3rem;
            margin-bottom: 5px; 
            font-size: 4.5rem; 
            font-weight: 800;
            line-height: 1.2;
        }
        .hero-highlight {
            background: linear-gradient(90deg, #2dd4bf, #06b6d4);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
        }
        .hero-subtitle { 
            text-align: center; 
            color: #8892b0; 
            margin: 10px auto 40px auto; 
            font-size: 1.15rem;
            font-weight: 400;
            max-width: 650px;
        }

        /* Buttons */
        div.stButton > button[kind="primary"] {
            background-color: #2dd4bf !important;
            color: #0f1519 !important;
            border: none !important;
            font-weight: 600 !important;
            border-radius: 8px !important;
        }
        div.stButton > button[kind="secondary"] {
            background-color: rgba(255, 255, 255, 0.05) !important;
            color: #ffffff !important;
            border: 1px solid rgba(255, 255, 255, 0.1) !important;
            font-weight: 600 !important;
            border-radius: 8px !important;
        }

        /* File Uploader Dropzone mimicking a large rectangle */
        [data-testid="stFileUploadDropzone"] {
            border: 2px dashed rgba(45, 212, 191, 0.4) !important;
            background: rgba(45, 212, 191, 0.02) !important;
            border-radius: 12px !important;
            padding: 50px !important;
            min-height: 240px;
            display: flex;
            align-items: center;
            justify-content: center;
            transition: all 0.3s ease;
        }
        [data-testid="stFileUploadDropzone"]:hover {
            border: 2px dashed rgba(45, 212, 191, 0.8) !important;
            background: rgba(45, 212, 191, 0.05) !important;
        }

        /* Miscellaneous Cards */
        .glass-panel { padding: 20px; max-width: 500px; margin: auto; }

        .metric-card {
            background: rgba(255, 255, 255, 0.03);
            padding: 20px;
            border-radius: 12px;
            text-align: center;
            border-left: 4px solid #2dd4bf;
            box-shadow: 0 4px 6px rgba(0,0,0,0.1);
        }

        .metric-label { font-size: 0.8rem; text-transform: uppercase; color: rgba(255, 255, 255, 0.5); font-weight: 600;}
        .metric-value { font-size: 1.8rem; font-weight: 800; color: #ffffff; margin-top: 5px;}
        .metric-caption { font-size: 0.75rem; color: #2dd4bf; margin-top: 4px; font-weight: 600; }

        .sidebar-title { 
            color: #2dd4bf; font-weight: 600; font-size: 1.3rem; 
            border-bottom: 1px solid rgba(45, 212, 191, 0.2);
            padding-bottom: 10px; margin-bottom: 20px;
        }

        .account-box {
            background: rgba(45, 212, 191, 0.05);
            padding: 15px; border-radius: 10px;
            border: 1px solid rgba(45, 212, 191, 0.2);
        }

        .word-pill {
            display: inline-block;
            background: rgba(45, 212, 191, 0.1);
            color: #2dd4bf;
            padding: 5px 12px;
            border-radius: 20px;
            margin: 4px;
            font-size: 0.85rem;
            border: 1px solid rgba(45, 212, 191, 0.3);
            font-weight: 500;
        }

        .role-tag {
            font-size: 0.7rem;
            background: #2dd4bf;
            color: #0f1519;
            padding: 2px 8px;
            border-radius: 4px;
            font-weight: bold;
            text-transform: uppercase;
        }

        .export-section {
            background: rgba(45, 212, 191, 0.03);
            border: 1px dashed rgba(45, 212, 191, 0.2);
            border-radius: 12px;
            padding: 25px;
            text-align: center;
            margin-top: 30px;
        }
        
        .summary-box {
            background: rgba(45, 212, 191, 0.03);
            border: 1px solid rgba(45, 212, 191, 0.15);
            border-radius: 10px;
            padding: 20px;
            margin-bottom: 15px;
            transition: all 0.3s ease;
        }
        
        .summary-box:hover {
            background: rgba(45, 212, 191, 0.08);
            border-color: rgba(45, 212, 191, 0.4);
        }
        
        .summary-title {
            color: #2dd4bf;
            font-weight: 600;
            margin-bottom: 10px;
            font-size: 1.15rem;
            border-bottom: 1px dashed rgba(45, 212, 191, 0.2);
            padding-bottom: 8px;
        }
        
        .summary-content {
            color: #d1d5db;
            line-height: 1.6;
            font-size: 0.95rem;
        }

        /* Hide default Streamlit elements that clutter */
        #MainMenu {visibility: hidden;}
        footer {visibility: hidden;}

    </style>
""", unsafe_allow_html=True)

# -----------------------------
# Core Logic Functions
# -----------------------------

def get_frequent_keywords(text, n=12):
    stop_words = set(['the', 'and', 'that', 'with', 'from', 'this', 'their', 'which', 'were', 'also', 'been', 'used', 'these', 'could', 'would'])
    words = re.findall(r'\b[a-z]{4,}\b', text.lower())
    filtered = [w for w in words if w not in stop_words]
    return Counter(filtered).most_common(n)

def extractive_summarize(text, num_sentences=3):
    if not text or len(text.split()) < 50: return text
    blob = TextBlob(text)
    sentences = blob.sentences
    if len(sentences) <= num_sentences: return text
    
    word_freq = {}
    for word in blob.words:
        w = word.lower()
        if w.isalpha():
            word_freq[w] = word_freq.get(w, 0) + 1
            
    max_f = max(word_freq.values()) if word_freq else 1
    for w in word_freq: word_freq[w] /= max_f
    
    sent_scores = {}
    for sent in sentences:
        for word in sent.words:
            w = word.lower()
            if w in word_freq:
                sent_scores[sent] = sent_scores.get(sent, 0) + word_freq[w]
                
    top_sents = heapq.nlargest(num_sentences, sent_scores, key=sent_scores.get)
    return " ".join([str(s) for s in sorted(top_sents, key=lambda x: sentences.index(x))])

def calculate_readability(text):
    words = len(text.split())
    sents = text.count('.') + text.count('?') + 1
    if sents == 0 or words == 0: return 0
    score = 206.835 - 1.015 * (words / sents) - 84.6 * (1.5) 
    return max(0, min(100, score))

def analyze_linguistics(text):
    blob = TextBlob(text)
    words = blob.words
    sents = blob.sentences
    word_count = len(words)
    if word_count == 0: return None
    
    avg_sent_len = np.mean([len(s.words) for s in sents])
    avg_word_len = np.mean([len(w) for w in words]) if word_count > 0 else 0
    sentiment = blob.sentiment.polarity
    
    lang_score = min(100, (avg_sent_len * 1.2) + (50 + sentiment * 20))
    trans = ["however", "therefore", "thus", "consequently", "moreover"]
    coherence = min(100, (sum(text.lower().count(t) for t in trans) * 5) + 40)
    
    reasoning_keys = ["because", "since", "implies", "evidence", "result"]
    reasoning = min(100, (sum(text.lower().count(k) for k in reasoning_keys) * 4) + 35)
    
    sophistication = min(100, (len([w for w in words if len(w) > 7]) / word_count) * 400)
    readability = calculate_readability(text)
    
    composite = (lang_score * 0.3) + (coherence * 0.2) + (reasoning * 0.2) + (sophistication * 0.15) + (readability * 0.15)
    
    return {
        "metrics": {
            "Language": round(lang_score, 1),
            "Coherence": round(coherence, 1),
            "Reasoning": round(reasoning, 1),
            "Sophistication": round(sophistication, 1),
            "Readability": round(readability, 1),
            "Composite": round(composite, 1)
        },
        "stats": {
            "Word Count": word_count,
            "Sentences": len(sents),
            "Avg Sent Len": round(avg_sent_len, 2),
            "Avg Word Len": round(avg_word_len, 2),
            "Sentiment": round(sentiment, 2)
        },
        "long_sentences": [str(s) for s in sents if len(s.words) > 35][:5]
    }

def extract_sections_structured(file):
    data = {"Executive Summary": []}
    curr = "Executive Summary"
    text_content = ""
    
    if file.name.endswith(".pdf"):
        with pdfplumber.open(file) as pdf:
            for page in pdf.pages:
                lines = (page.extract_text() or "").split('\n')
                for line in lines:
                    clean = line.strip()
                    if not clean: continue
                    if (clean.isupper() and len(clean) < 60) or re.match(r'^\d+\.', clean):
                        curr = clean
                        data[curr] = []
                    else:
                        data[curr].append(clean)
                    text_content += clean + " "
    elif file.name.endswith(".docx"):
        doc = docx.Document(file)
        for para in doc.paragraphs:
            clean = para.text.strip()
            if not clean: continue
            if para.style.name.startswith('Heading') or clean.isupper():
                curr = clean
                data[curr] = []
            else:
                data[curr].append(clean)
            text_content += clean + " "
    else:
        text_content = file.getvalue().decode()
        data["Full Content"] = [text_content]

    raw_sections = {k: " ".join(v) for k, v in data.items() if v}
    return raw_sections, text_content

def create_pdf_report(results, filename):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, f"PaperIQ Analysis: {filename}", ln=True, align='C')
    pdf.ln(10)
    
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 10, f"Composite Score: {results['metrics']['Composite']}/100", ln=True)
    pdf.ln(2)
    
    pdf.set_font("Arial", '', 10)
    for k, v in results['metrics'].items():
        if k != 'Composite': pdf.cell(0, 8, f"{k}: {v}/100", ln=True)
    pdf.ln(5)

    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 10, "Text Statistics:", ln=True)
    pdf.ln(2)
    
    pdf.set_font("Arial", '', 10)
    for k, v in results['stats'].items():
        pdf.cell(0, 8, f"{k}: {v}", ln=True)
    pdf.ln(5)

    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 10, "Section Breakdowns:", ln=True)
    pdf.ln(3)
    
    for title, summary in results['summaries'].items():
        pdf.set_font("Arial", 'B', 10)
        safe_title = title.encode('latin-1', 'replace').decode('latin-1')
        pdf.cell(0, 8, f"-> {safe_title}", ln=True)
        
        pdf.set_font("Arial", '', 10)
        safe_summary = summary.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 5, safe_summary)
        pdf.ln(4)
        
    return pdf.output(dest='S').encode('latin-1', 'replace')

# -----------------------------
# UI Components
# -----------------------------

def auth_page():
    st.markdown("""
        <div style="display: flex; flex-direction: column; align-items: center; text-align: center; width: 100%;">
            <h1 class="hero-title"><span class="hero-highlight">PaperIQ</span></h1>
            <p class="hero-subtitle">AI Powered Research Insight Analyzer</p>
        </div>
    """, unsafe_allow_html=True)
    
    st.markdown("<br>", unsafe_allow_html=True)
    
    # Use columns to perfectly center the login panel
    _, center_col, _ = st.columns([1, 1.5, 1])
    
    with center_col:
        if st.session_state.auth_mode == "login":
            user = st.text_input("Username")
            pwd = st.text_input("Password", type="password")
            if st.button("Sign In", use_container_width=True, type="primary"):
                res = verify_user(user, pwd)
                if res:
                    st.session_state.logged_in = True
                    st.session_state.user_data = {"username": user, "gmail": res[0], "role": res[1]}
                    st.rerun()
                else:
                    st.error("Invalid credentials.")
            if st.button("Need an account? Sign Up", use_container_width=True):
                st.session_state.auth_mode = "signup"
                st.rerun()
        
        else:
            new_user = st.text_input("Choose Username")
            new_mail = st.text_input("Gmail Address")
            new_pwd = st.text_input("Password", type="password")
            new_role = st.selectbox("Your Role", ["Student", "Teacher", "Researcher", "Professional"])
            if st.button("Create Account", use_container_width=True, type="primary"):
                if new_user and new_mail and new_pwd:
                    if add_user(new_user, new_mail, new_pwd, new_role):
                        st.success("Account created! Please login.")
                        st.session_state.auth_mode = "login"
                        st.rerun()
                    else:
                        st.error("Username already exists.")
                else:
                    st.warning("All fields are required.")
            if st.button("Back to Login", use_container_width=True):
                st.session_state.auth_mode = "login"
                st.rerun()

def main_dashboard():
    with st.sidebar:
        st.markdown('<p style="font-weight:bold; font-size:1.5rem; color:#ffffff; margin-bottom: 20px;">✨ PaperIQ</p>', unsafe_allow_html=True)
        st.markdown('<p class="sidebar-title">Profile</p>', unsafe_allow_html=True)
        st.markdown(f"""<div class="account-box">
            <span class="role-tag">{st.session_state.user_data['role']}</span>
            <p style="margin:5px 0 0 0; font-weight:bold; color:#2dd4bf; font-size:1.1rem;">{st.session_state.user_data['username']}</p>
            <p style="margin:0; font-size:0.8rem; opacity:0.7;">{st.session_state.user_data['gmail']}</p>
        </div>""", unsafe_allow_html=True)
        
        st.markdown('<p class="sidebar-title" style="margin-top:30px;">History</p>', unsafe_allow_html=True)
        
        history_rows = get_history(st.session_state.user_data['username'])
        if history_rows:
            for row_id, fname, date_str, data_json in history_rows:
                short_fname = fname if len(fname) < 20 else fname[:17] + "..."
                col1, col2 = st.columns([5, 1])
                with col1:
                    if st.button(f"📄 {short_fname}", key=f"hist_{row_id}", help=f"Analyzed on: {date_str}", use_container_width=True):
                        st.session_state.analysis_results = json.loads(data_json)
                        st.rerun()
                with col2:
                    with st.popover("⋮"):
                        st.write("Options")
                        if st.button("🗑️ Delete", key=f"del_{row_id}", type="primary"):
                            delete_history(row_id)
                            st.rerun()
        else:
            st.caption("No past analyses found.")
            
        st.divider()
        if st.button("Logout", use_container_width=True):
            st.session_state.logged_in = False
            st.session_state.analysis_results = None
            st.rerun()

    # Main Dashboard Hero
    st.markdown("""
        <div style="display: flex; flex-direction: column; align-items: center; text-align: center; width: 100%;">
            <h1 class="hero-title"><span class="hero-highlight">PaperIQ</span></h1>
            <p class="hero-subtitle">Upload a PDF, DOCX, or TXT file to get deep insights, structured summaries, and quality metrics in seconds.</p>
        </div>
    """, unsafe_allow_html=True)
    
    st.markdown("<br>", unsafe_allow_html=True)
    
    # File Uploader Area (Centered and wider)
    c_pad1, c_main, c_pad2 = st.columns([1, 3, 1])
    with c_main:
        uploaded_file = st.file_uploader("Drop your research paper here", type=["pdf", "docx", "txt"], label_visibility="collapsed")
        if uploaded_file:
            if st.button("Analyze Document", type="primary", use_container_width=True):
                with st.spinner("Extracting intelligence..."):
                    raw_sections, full_text = extract_sections_structured(uploaded_file)
                    data = analyze_linguistics(full_text)
                    keywords = get_frequent_keywords(full_text)
                    results = {
                        "raw_sections": raw_sections,
                        "summaries": {}, 
                        "metrics": data["metrics"],
                        "stats": data["stats"],
                        "issues": data["long_sentences"],
                        "keywords": keywords,
                        "filename": uploaded_file.name,
                        "raw_text": full_text
                    }
                    st.session_state.analysis_results = results
                    save_history(st.session_state.user_data['username'], uploaded_file.name, results)
                    st.rerun()

    st.markdown('<p style="text-align: center; color: rgba(255,255,255,0.4); font-size: 0.8rem; margin-top: 20px;">AI-generated insights. Always refer to the original paper for authoritative information.</p>', unsafe_allow_html=True)

    # Dashboard Rendering
    if st.session_state.analysis_results:
        st.divider()
        res = st.session_state.analysis_results
        
        st.markdown(f"<div style='text-align: center; color: #2dd4bf; margin-bottom: 20px;'><strong>Currently Viewing Analysis for:</strong> {res['filename']}</div>", unsafe_allow_html=True)

        s_val = res["stats"]["Sentiment"]
        s_label = "Positive / Assertive" if s_val > 0.1 else "Critical / Negative" if s_val < -0.1 else "Neutral / Objective"

        m1, m2, m3, m4 = st.columns(4)
        m1.markdown(f'<div class="metric-card"><div class="metric-label">Composite</div><div class="metric-value">{res["metrics"]["Composite"]}</div></div>', unsafe_allow_html=True)
        m2.markdown(f'<div class="metric-card"><div class="metric-label">Words</div><div class="metric-value">{res["stats"]["Word Count"]:,}</div></div>', unsafe_allow_html=True)
        m3.markdown(f'<div class="metric-card"><div class="metric-label">Sentences</div><div class="metric-value">{res["stats"]["Sentences"]}</div></div>', unsafe_allow_html=True)
        m4.markdown(f'<div class="metric-card"><div class="metric-label">Sentiment</div><div class="metric-value">{s_val}</div><div class="metric-caption">{s_label}</div></div>', unsafe_allow_html=True)

        st.markdown("### 🏷️ Lexical Keywords")
        kw_html = "".join([f'<span class="word-pill">{w} ({c})</span>' for w, c in res["keywords"]])
        st.markdown(f'<div>{kw_html}</div>', unsafe_allow_html=True)
        
        st.markdown("<br>", unsafe_allow_html=True)

        t1, t2, t3, t4, t5, t6 = st.tabs([
            "📊 Visualisations", 
            "📑 Section Breakdowns", 
            "⚠️ Complexity Alerts",
            "💡 Suggestions", 
            "📝 Detailed Metrics", 
            "❤️ Sentiment"
        ])
        
        with t1:
            col_a, col_b = st.columns(2)
            with col_a:
                st.markdown("#### Metric Radar")
                cats = list(res["metrics"].keys())[:-1]
                vals = [res["metrics"][c] for c in cats]
                fig = go.Figure(data=[go.Scatterpolar(r=vals, theta=cats, fill='toself', line_color='#2dd4bf')])
                fig.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 100])), paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font_color='white', height=380)
                st.plotly_chart(fig, use_container_width=True)
            with col_b:
                st.markdown("#### Text Statistics")
                fig_bar = go.Figure(data=[
                    go.Bar(name='Avg Sentence Len', x=['Avg Sentence Len'], y=[res['stats']['Avg Sent Len']], marker_color='#2dd4bf'),
                    go.Bar(name='Avg Word Len', x=['Avg Word Len'], y=[res['stats']['Avg Word Len']], marker_color='#06b6d4')
                ])
                fig_bar.update_layout(
                    paper_bgcolor='rgba(0,0,0,0)',
                    plot_bgcolor='rgba(0,0,0,0)',
                    font_color='white',
                    height=380,
                    margin=dict(t=20, b=20)
                )
                st.plotly_chart(fig_bar, use_container_width=True)
                
            st.markdown("#### Detailed Scores")
            st.markdown(f"""
            <div style="color: #2dd4bf; padding: 15px; background: rgba(45, 212, 191, 0.05); border-radius: 10px; border: 1px solid rgba(45, 212, 191, 0.2);">
                <div style="display: flex; justify-content: space-around; flex-wrap: wrap;">
                    <p style="margin: 0; line-height: 1.8;"><strong>Language:</strong> {res['metrics']['Language']}/100</p>
                    <p style="margin: 0; line-height: 1.8;"><strong>Coherence:</strong> {res['metrics']['Coherence']}/100</p>
                    <p style="margin: 0; line-height: 1.8;"><strong>Reasoning:</strong> {res['metrics']['Reasoning']}/100</p>
                    <p style="margin: 0; line-height: 1.8;"><strong>Sophistication:</strong> {res['metrics']['Sophistication']}/100</p>
                    <p style="margin: 0; line-height: 1.8;"><strong>Readability:</strong> {res['metrics']['Readability']}/100</p>
                </div>
                <hr style="border-color: rgba(45, 212, 191, 0.2); margin: 12px 0;">
                <p style="text-align: center; margin: 0;"><strong>Overall Tone:</strong> {s_label}</p>
            </div>
            """, unsafe_allow_html=True)

        with t2:
            st.markdown("#### Extracted Section Summaries")
            summary_length = st.select_slider(
                "Adjust Summary Detail Level:", 
                options=["Short", "Medium", "Long"], 
                key="summary_length_slider" 
            )
            length_map = {"Short": 2, "Medium": 4, "Long": 7}
            num_sentences = length_map.get(summary_length, 4)
            res["summaries"] = {}
            for title, content in res["raw_sections"].items():
                summary = extractive_summarize(content, num_sentences)
                res["summaries"][title] = summary
                with st.expander(f"● {title}"):
                    st.write(summary)

        with t3:
            if res["issues"]:
                st.info("The following sentences may be too complex for general readers:")
                for issue in res["issues"]:
                    st.warning(f"“...{issue[:160]}...”")
            else:
                st.success("Document structure is clear and concise.")
                
        with t4:
            st.subheader("Vocabulary Improvements")
            suggestions_map = {"very": "extremely", "bad": "adverse", "good": "beneficial", "show": "demonstrate", "big": "substantial"}
            text_lower = res['raw_text'].lower()
            found = False
            for s, c in suggestions_map.items():
                if re.search(rf'\b{s}\b', text_lower):
                    st.info(f"Consider replacing **'{s}'** with **'{c}'** for a more academic tone.")
                    found = True
            if not found: st.success("Great vocabulary! No basic academic replacements suggested.")

        with t5:
            d_col1, d_col2 = st.columns(2)
            with d_col1:
                st.write(f"**Total Words:** {res['stats']['Word Count']}")
                st.write(f"**Total Sentences:** {res['stats']['Sentences']}")
            with d_col2:
                st.write(f"**Avg Sentence Length:** {res['stats']['Avg Sent Len']} words")
                st.write(f"**Avg Word Length:** {res['stats']['Avg Word Len']} characters")
                st.write(f"**Sophistication Score:** {res['metrics']['Sophistication']}%")

        with t6:
            st.metric("Sentiment Score", res['stats']['Sentiment'])
            if res['stats']['Sentiment'] > 0.1: 
                st.success("Positive Tone / Assertive")
            elif res['stats']['Sentiment'] < -0.1: 
                st.warning("Critical Tone / Negative")
            else: 
                st.info("Neutral Tone / Objective")

        # Export Section
        st.markdown('<div class="export-section">', unsafe_allow_html=True)
        st.markdown("### 📥 Document Intelligence Report")
        st.markdown("Download a PDF snapshot of the analysis including all linguistic scores.")
        pdf_data = create_pdf_report(res, res["filename"])
        st.download_button("Download Analysis PDF", data=pdf_data, file_name=f"PaperIQ_{res['filename']}.pdf", mime="application/pdf", use_container_width=True)
        st.markdown('</div>', unsafe_allow_html=True)

# -----------------------------
# Entry Point
# -----------------------------
if not st.session_state.logged_in:
    auth_page()
else:
    main_dashboard()
